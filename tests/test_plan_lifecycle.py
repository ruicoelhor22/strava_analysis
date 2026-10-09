"""Synthetic lifecycle tests; no private workbook or account needed."""
import json
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook

from tests.test_plan import _write_plan, _test_dir
from tests.test_calendar import FakeCalendarClient
from tests.test_coaching import _database, _plan_session, _activity
from endurance_lab.db import connect, init_db
from endurance_lab.plan_importer import import_training_plan
from endurance_lab.plan_revisions import sync_plan_revision
from endurance_lab.plan_lifecycle import revision_status, phase_context, lifecycle_rows, confirm_session
from endurance_lab.plan_adherence import planned_sessions
from endurance_lab.training_decisions import choose_training, save_training_decision, training_decisions
from endurance_lab.weekly_adjustment import adjust_week
from endurance_lab.google_calendar import sync_google_calendar
from endurance_lab.plan_matching import match_training_plan, set_manual_match
from endurance_lab.coaching_backtest import backtest


def fixture():
    root=_test_dir(); workbook=_write_plan(root/'coaching.xlsx'); database=root/'plan.sqlite3'
    import_training_plan(workbook,database)
    return root,workbook,database


def edit(workbook, cell, value):
    copy=load_workbook(workbook);copy['Training Calendar'][cell]=value;copy.save(workbook);copy.close()


def first_id(database):
    with connect(database) as connection:
        return connection.execute("SELECT id FROM planned_sessions WHERE sport='cycling' ORDER BY id LIMIT 1").fetchone()[0]


def test_unchanged_workbook_creates_no_revision_or_prescription_churn():
    _,workbook,database=fixture()
    before=revision_status(database)
    result=sync_plan_revision(database)
    after=revision_status(database)
    assert result['status']=='unchanged'
    assert before==after
    assert import_training_plan(workbook,database).unchanged


def test_duration_edit_is_material_and_keeps_identity_and_flags_choice():
    _,workbook,database=fixture();sid=first_id(database)
    choose_training(sid,'KEEP_ORIGINAL',database)
    edit(workbook,'K4',120)
    update=sync_plan_revision(database)
    assert update['status']=='imported'
    assert first_id(database)==sid
    change=next(c for c in update['changes'] if c['type']=='DURATION_CHANGED')
    assert change['before']==3600 and change['after']==7200 and change['material']
    assert training_decisions(database)[sid]['needs_review']==1
    assert training_decisions(database)[sid]['workout']['duration_seconds']==3600
    assert revision_status(database)['decision_conflicts']==1
    choose_training(sid,'KEEP_ORIGINAL',database,note='Reviewed updated plan')
    assert training_decisions(database)[sid]['workout']['duration_seconds']==7200
    assert not training_decisions(database)[sid]['needs_review']


def test_typo_is_non_material_and_unrelated_sessions_unchanged():
    _,workbook,database=fixture();sid=first_id(database)
    choose_training(sid,'KEEP_ORIGINAL',database)
    edit(workbook,'I4','Threshold 3x10!')
    summary=import_training_plan(workbook,database)
    assert summary.changes and all(not c['material'] for c in summary.changes)
    assert first_id(database)==sid
    assert not training_decisions(database)[sid]['needs_review']
    assert len({c.get('session_id') for c in summary.changes})==1


def test_move_and_type_edit_preserve_defensible_identity():
    _,workbook,database=fixture();sid=first_id(database)
    edit(workbook,'A4',datetime(2026,9,23))
    summary=import_training_plan(workbook,database)
    assert first_id(database)==sid
    assert any(c['type']=='SESSION_MOVED' for c in summary.changes)
    edit(workbook,'G4','Run')
    summary=import_training_plan(workbook,database)
    assert any(c['type']=='SESSION_TYPE_CHANGED' and c['session_id']==sid for c in summary.changes)


def test_removed_future_session_is_retired_not_deleted():
    _,workbook,database=fixture();sid=first_id(database)
    for cell in ('G4','I4','J4','K4'):edit(workbook,cell,None)
    summary=import_training_plan(workbook,database)
    assert any(c['type']=='SESSION_REMOVED' and c['session_id']==sid for c in summary.changes)
    with connect(database) as connection:
        assert connection.execute('SELECT active FROM planned_sessions WHERE id=?',(sid,)).fetchone()[0]==0
        assert connection.execute('SELECT count(*) FROM plan_session_versions WHERE planned_session_id=?',(sid,)).fetchone()[0]>=1


def test_completed_historical_workout_and_match_survive_coach_edits():
    _,workbook,database=fixture();sid=first_id(database)
    actual=_activity(database,'2026-09-21','cycling',3600)
    set_manual_match(sid,actual,database=database)
    edit(workbook,'K4',120)
    summary=import_training_plan(workbook,database)
    match_training_plan(database)
    with connect(database) as connection:
        assert connection.execute('SELECT planned_duration_seconds FROM planned_sessions WHERE id=?',(sid,)).fetchone()[0]==3600
        assert connection.execute('SELECT activity_id FROM planned_activity_matches WHERE planned_session_id=? AND is_selected=1',(sid,)).fetchone()[0]==actual
    assert any(c.get('completed_match_preserved') for c in summary.changes)


def test_invalid_revision_and_dry_run_cannot_replace_valid_plan():
    _,workbook,database=fixture();before=revision_status(database)
    edit(workbook,'K4',120)
    preview=import_training_plan(workbook,database,dry_run=True)
    assert preview.changes
    assert revision_status(database)==before
    copy=load_workbook(workbook);copy['Training Calendar'].title='Invalid';copy.save(workbook);copy.close()
    result=sync_plan_revision(database)
    assert result['status']=='failed' and result['existing_plan_preserved']
    assert revision_status(database)==before


def test_draft_incomplete_and_blank_days_are_not_confirmed():
    _,workbook,database=fixture()
    edit(workbook,'S4','Draft')
    edit(workbook,'K5',None)
    import_training_plan(workbook,database)
    states=lifecycle_rows(database)
    assert states[first_id(database)]['state']=='draft'
    assert 'incomplete' in {r['state'] for r in states.values()}
    assert len(revision_status(database)['revision']['missing_dates'])==5
    assert all(p.prescribed is None for p in adjust_week(date(2026,9,21),database,as_of=date(2026,9,21),persist=False))
    confirm_session(first_id(database),database)
    assert lifecycle_rows(database)[first_id(database)]['state']=='confirmed'


@pytest.mark.parametrize('action',['ACCEPT','KEEP_ORIGINAL','MODIFY','MOVE','SKIP','CLEAR'])
def test_decision_actions_preserve_independent_app_prescription(action):
    database=_database();sid=_plan_session(database,'2026-10-02','cycling','Easy Z2',3600)
    adjust_week(date(2026,10,1),database,as_of=date(2026,10,1))
    workout={'date':'2026-10-03' if action=='MOVE' else '2026-10-02','sport':'cycling','title':'My Z2','duration_seconds':2700}
    if action=='CLEAR':choose_training(sid,'KEEP_ORIGINAL',database)
    choose_training(sid,action,database,workout=None if action=='ACCEPT' else workout)
    decisions=training_decisions(database)
    if action=='CLEAR': assert sid not in decisions
    else: assert decisions[sid]['action']==action
    with connect(database) as connection:
        app=connection.execute("SELECT * FROM workout_prescriptions WHERE planned_session_id=? AND status='active'",(sid,)).fetchone()
        assert app['action']=='KEEP' and json.loads(app['prescribed_json'])['duration_seconds']==3600
        assert connection.execute('SELECT count(*) FROM athlete_decision_history WHERE planned_session_id=?',(sid,)).fetchone()[0]>=1


def test_calendar_effective_move_skip_draft_and_idempotency():
    root,workbook,database=fixture();sid=first_id(database)
    config=root/'calendar.yaml'
    config.write_text('google_calendar:\n  enabled: true\n  planned_calendar_id: test-planned\n',encoding='utf-8')
    client=FakeCalendarClient();today=date(2026,9,20)
    first=sync_google_calendar(database,client=client,today=today,config_path=config)
    assert first.created==2
    workout={'date':'2026-09-23','sport':'cycling','title':'Moved threshold','duration_seconds':3600}
    choose_training(sid,'MOVE',database,workout=workout)
    moved=sync_google_calendar(database,client=client,today=today,config_path=config)
    assert moved.updated==1 and moved.created==0
    assert client.updated[-1][1]['start']['date']=='2026-09-23'
    assert sync_google_calendar(database,client=client,today=today,config_path=config).unchanged==2
    choose_training(sid,'SKIP',database)
    assert sync_google_calendar(database,client=client,today=today,config_path=config).removed==1
    choose_training(sid,'CLEAR',database)
    edit(workbook,'S4','Draft');import_training_plan(workbook,database)
    assert sync_google_calendar(database,client=client,today=today,config_path=config).considered==1


def test_accept_cannot_silently_replace_saved_rest_with_gym():
    database=_database();sid=_plan_session(database,'2026-10-01','strength','Gym B',3000)
    rest={'date':'2026-10-01','sport':'other','title':'Rest','duration_seconds':0}
    save_training_decision(sid,rest,'Rest today; gym Friday',database)
    adjust_week(date(2026,10,1),database,as_of=date(2026,10,1))
    with pytest.raises(ValueError,match='replace your saved decision'):
        choose_training(sid,'ACCEPT',database)
    assert training_decisions(database)[sid]['workout']==rest
    choose_training(sid,'RETAIN',database)
    assert training_decisions(database)[sid]['workout']==rest
    choose_training(sid,'ACCEPT',database,replace_decision=True)
    assert training_decisions(database)[sid]['workout']['sport']=='strength'


def test_retain_an_accepted_choice_does_not_reaccept_changed_app_recommendation():
    database=_database();sid=_plan_session(database,'2026-10-02','cycling','Easy Z2',3600)
    adjust_week(date(2026,10,1),database,as_of=date(2026,10,1))
    choose_training(sid,'ACCEPT',database)
    accepted=training_decisions(database)[sid]['workout']
    changed={**accepted,'duration_seconds':1800,'title':'New app recommendation'}
    with connect(database) as connection:
        connection.execute("UPDATE workout_prescriptions SET prescribed_json=? WHERE planned_session_id=? AND status='active'",(json.dumps(changed),sid))
        connection.commit()
    choose_training(sid,'RETAIN',database,note='Keep my previously accepted workout')
    assert training_decisions(database)[sid]['workout']==accepted
    with pytest.raises(ValueError,match='recommendation changed'):
        choose_training(sid,'ACCEPT',database,workout=accepted,replace_decision=True)


def test_reordered_long_ride_has_endurance_candidate_not_threshold_credit():
    database=_database()
    threshold=_plan_session(database,'2026-10-06','cycling','Threshold touch',4500,intensity='Threshold')
    long_ride=_plan_session(database,'2026-10-10','cycling','Saturday endurance ride',9900,intensity='Z2')
    actual=_activity(database,'2026-10-05','cycling',12000,load=210,intensity_factor=0.79)
    with connect(database) as connection:
        connection.execute('UPDATE derived_activity_metrics SET power_zones_json=? WHERE activity_id=?',(json.dumps([5600,2400,1500,1100,700,700]),actual))
        connection.commit()
    match_training_plan(database)
    with connect(database) as connection:
        candidate=connection.execute("SELECT planned_session_id,is_selected,match_status FROM planned_activity_matches WHERE activity_id=?",(actual,)).fetchone()
        assert tuple(candidate)==(long_ride,0,'probable')
        assert not connection.execute('SELECT 1 FROM planned_activity_matches WHERE planned_session_id=?',(threshold,)).fetchone()


def test_gym_reordered_one_day_late_links_to_moved_workout_not_other_sport():
    database=_database()
    gym=_plan_session(database,'2026-10-05','strength','Gym A deload',2400)
    _plan_session(database,'2026-10-06','cycling','Threshold touch',4500,intensity='Threshold')
    choose_training(gym,'MOVE',database,workout={'date':'2026-10-06','sport':'strength','title':'Gym A deload','duration_seconds':2400})
    actual=_activity(database,'2026-10-06','strength',2400)
    match_training_plan(database)
    with connect(database) as connection:
        selected=connection.execute('SELECT planned_session_id,activity_id FROM planned_activity_matches WHERE is_selected=1').fetchone()
        assert tuple(selected)==(gym,actual)


def test_skipped_key_does_not_create_fake_post_key_recovery_gate():
    database=_database()
    key=_plan_session(database,'2026-10-06','cycling','Threshold touch',4500,priority='A',intensity='Threshold')
    easy=_plan_session(database,'2026-10-07','cycling','Easy Z2',4200,priority='C',intensity='Easy aerobic')
    choose_training(key,'SKIP',database)
    result=next(p for p in adjust_week(date(2026,10,6),database,as_of=date(2026,10,6),persist=False) if p.planned_session_id==easy)
    assert 'protect_key_session_recovery' not in result.rules_triggered


def test_confirmed_extra_endurance_stays_extra_across_repeated_matching():
    from endurance_lab.plan_matching import mark_extra_activity
    from endurance_lab.plan_adherence import actual_training_days
    database=_database()
    _plan_session(database,'2026-10-10','cycling','Saturday endurance ride',9900,intensity='Z2')
    actual=_activity(database,'2026-10-05','cycling',12000,load=210,intensity_factor=0.79)
    with connect(database) as connection:
        connection.execute('UPDATE derived_activity_metrics SET power_zones_json=? WHERE activity_id=?',(json.dumps([5600,2400,1500,1100,700,700]),actual));connection.commit()
    match_training_plan(database)
    mark_extra_activity(actual,'Athlete confirms extra holiday endurance, not Saturday replacement',database)
    match_training_plan(database);match_training_plan(database)
    with connect(database) as connection:
        assert not connection.execute('SELECT 1 FROM planned_activity_matches WHERE activity_id=?',(actual,)).fetchone()
    days=actual_training_days(database,start=date(2026,10,5),end=date(2026,10,11))
    assert days['2026-10-05'][0]['role']=='extra'
    assert days['2026-10-05'][0]['duration_seconds']==12000


@pytest.mark.parametrize('muscular,confirmed_rest,expected',[('moderate',True,'REDUCE_INTENSITY'),('moderate',False,'REPLACE_WITH_EASY'),('high',True,'REPLACE_WITH_EASY')])
def test_confirmed_rest_allows_only_moderate_deload_strength(monkeypatch,muscular,confirmed_rest,expected):
    from endurance_lab.coaching_models import SessionCost
    import endurance_lab.weekly_adjustment as engine
    database=_database()
    gym=_plan_session(database,'2026-10-08','strength','Gym B deload',2400,intensity='Deload strength')
    rest=_plan_session(database,'2026-10-07','cycling','Easy Z2',4200)
    if confirmed_rest:choose_training(rest,'MODIFY',database,workout={'date':'2026-10-07','sport':'other','title':'Rest','duration_seconds':0})
    cost=SessionCost(1,'strength_local','low','low',muscular,'full','strength','low','moderate','high')
    monkeypatch.setattr(engine,'_actual_context',lambda *a,**k:[{'id':1,'sport':'strength','started_at':datetime(2026,10,6,17,tzinfo=timezone.utc),'cost':cost}])
    result=next(p for p in adjust_week(date(2026,10,8),database,as_of=date(2026,10,8),persist=False) if p.planned_session_id==gym)
    assert result.action==expected
    if confirmed_rest and muscular=='moderate':
        assert result.prescribed['sport']=='strength'
        assert 'confirmed_rest_between_moderate_strength_sessions' in result.rules_triggered


def test_day_card_shows_completed_extra_even_with_no_planned_workout(monkeypatch):
    from dashboard import plan_pages
    from datetime import timedelta
    import streamlit as st
    captured=[]
    monkeypatch.setattr(st,'markdown',lambda body,**kwargs:captured.append(body))
    monday=date(2026,10,5)
    plan_pages._week_grid([],monday,monday+timedelta(days=1),actual_days={'2026-10-05':[{
        'id':7,'name':'Holiday endurance','sport':'cycling','role':'extra','linked_session_id':None,
        'duration_seconds':12000,'distance_m':99100,'metadata_only':0}]})
    html=captured[-1]
    monday_card=html.split('<header>Mon 05</header>')[1].split('<header>Tue 06</header>')[0]
    assert 'Holiday endurance' in monday_card and 'Extra training' in monday_card and 'Done' in monday_card
    assert 'No planned session' not in monday_card


def test_key_session_runway_uses_its_own_date_not_later_weekend_key():
    database=_database()
    sid=_plan_session(database,'2026-10-06','cycling','Threshold touch',3600,priority='A',intensity='Threshold')
    _plan_session(database,'2026-10-10','cycling','Long ride',7200,priority='A',intensity='Z2')
    _activity(database,'2026-10-05','cycling',10800,load=220,intensity_factor=0.8)
    result=next(p for p in adjust_week(date(2026,10,6),database,as_of=date(2026,10,6),persist=False) if p.planned_session_id==sid)
    assert result.recovery_runway_h is not None and result.recovery_runway_h < 48


def test_saved_rest_ui_defaults_to_keep_and_requires_explicit_replacement(monkeypatch):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    from dashboard import plan_controls
    database=_database();sid=_plan_session(database,'2026-10-01','strength','Gym B',3000)
    rest={'date':'2026-10-01','sport':'other','title':'Rest','duration_seconds':0}
    save_training_decision(sid,rest,'Rest today',database)
    adjust_week(date(2026,10,1),database,as_of=date(2026,10,1))
    monkeypatch.setenv('ENDURANCE_DB_PATH',str(database.resolve()))
    monkeypatch.setattr(plan_controls,'_refresh',lambda:None)
    app=AppTest.from_file(str(Path(__file__).parent/'fixtures'/'plan_controls_app.py')).run()
    assert not app.exception
    assert app.selectbox[0].value=='Keep my decision'
    next(b for b in app.button if b.label=='Save decision').click().run()
    assert training_decisions(database)[sid]['workout']==rest
    app.selectbox[0].select('Accept app recommendation').run()
    next(b for b in app.button if b.label=='Save decision').click().run()
    assert not app.exception and app.error
    assert training_decisions(database)[sid]['workout']==rest
    app.checkbox[0].check().run()
    next(b for b in app.button if b.label=='Save decision').click().run()
    assert not app.exception
    assert training_decisions(database)[sid]['workout']['sport']=='strength'


@pytest.mark.parametrize(('text','expected'),[('Base','BASE'),('Threshold Build 2','BUILD'),('Holiday / reset','RECOVERY'),('Taper','TAPER'),('Race','RACE'),('Unclear block','UNKNOWN')])
def test_phase_mapping_preserves_original_and_week_context(text,expected):
    result=phase_context(text,'Deload')
    assert result['phase']==expected and result['phase_text']==text and result['week_type']=='RECOVERY'


@pytest.mark.parametrize('phase',['Base','Build','Recovery','Taper','Race','Unknown'])
def test_phase_context_never_adds_volume(phase):
    database=_database();sid=_plan_session(database,'2026-10-02','cycling','Easy Z2',1800)
    with connect(database) as connection:
        plan_id=connection.execute('SELECT plan_id FROM planned_sessions WHERE id=?',(sid,)).fetchone()[0]
        week=connection.execute("INSERT INTO training_weeks(plan_id,week_start,week_end,phase_name,week_type,raw_json,active) VALUES (?,'2026-09-28','2026-10-04',?,'Deload','{}',1)",(plan_id,phase)).lastrowid
        connection.execute('UPDATE planned_sessions SET week_id=? WHERE id=?',(week,sid));connection.commit()
    result=adjust_week(date(2026,10,1),database,as_of=date(2026,10,1),persist=False)[0]
    assert result.prescribed['duration_seconds']<=1800
    assert 'respect_planned_recovery_volume' in result.rules_triggered


def test_taper_omits_leg_support_before_race_and_keeps_race_priority():
    database=_database();_plan_session(database,'2026-10-02','strength','Leg strength',2400,priority='C')
    _plan_session(database,'2026-10-03','cycling','Race',7200,priority='B')
    with connect(database) as connection:
        plan_id=connection.execute('SELECT id FROM training_plans').fetchone()[0]
        week=connection.execute("INSERT INTO training_weeks(plan_id,week_start,week_end,phase_name,raw_json,active) VALUES (?,'2026-09-28','2026-10-04','Taper','{}',1)",(plan_id,)).lastrowid
        connection.execute('UPDATE planned_sessions SET week_id=?',(week,));connection.commit()
    results=adjust_week(date(2026,10,1),database,as_of=date(2026,10,1),persist=False)
    assert results[0].action=='SKIP' and 'avoid_leg_damage_before_event' in results[0].rules_triggered
    assert results[1].decision_trace['planned_session']['priority']=='A'


def test_history_does_not_use_future_revision_or_future_decision():
    _,workbook,database=fixture();sid=first_id(database)
    with connect(database) as connection:
        connection.execute("UPDATE plan_revisions SET imported_at='2026-09-01T00:00:00+00:00'");connection.commit()
    choose_training(sid,'KEEP_ORIGINAL',database)
    edit(workbook,'K4',120);import_training_plan(workbook,database)
    historical=adjust_week(date(2026,9,21),database,as_of=date(2026,9,21),persist=False,strict_history=True)
    assert historical[0].prescribed['duration_seconds']==3600
    assert historical[0].action!='USER_CHOICE'
    assert not backtest(date(2026,8,1),date(2026,8,7),database)['decisions']


def test_removed_workout_with_choice_stays_visible_as_a_review_conflict():
    _,workbook,database=fixture();sid=first_id(database)
    choose_training(sid,'KEEP_ORIGINAL',database)
    for cell in ('G4','I4','J4','K4'):edit(workbook,cell,None)
    import_training_plan(workbook,database)
    rows=planned_sessions(database)
    row=next(r for r in rows if r['id']==sid)
    assert row['needs_review'] and row['lifecycle']=='cancelled' and row['effective']['duration_seconds']==3600
    choose_training(sid,'MODIFY',database,workout=row['effective'],note='Retain my choice despite coach removal')
    assert not training_decisions(database)[sid]['needs_review']
    choose_training(sid,'CLEAR',database)
    assert sid not in {r['id'] for r in planned_sessions(database)}


def test_calendar_can_restore_skipped_event_without_updating_deleted_remote_event():
    root,_,database=fixture();sid=first_id(database);config=root/'calendar.yaml'
    config.write_text('google_calendar:\n  enabled: true\n  planned_calendar_id: test-planned\n',encoding='utf-8')
    client=FakeCalendarClient();today=date(2026,9,20)
    sync_google_calendar(database,client=client,today=today,config_path=config)
    choose_training(sid,'SKIP',database)
    sync_google_calendar(database,client=client,today=today,config_path=config)
    choose_training(sid,'CLEAR',database)
    restored=sync_google_calendar(database,client=client,today=today,config_path=config)
    assert restored.created==1 and restored.updated==0


def test_athlete_move_outside_horizon_updates_existing_event_identity():
    root,_,database=fixture();sid=first_id(database);config=root/'calendar.yaml'
    config.write_text('google_calendar:\n  enabled: true\n  planned_calendar_id: test-planned\n',encoding='utf-8')
    client=FakeCalendarClient();today=date(2026,9,20)
    sync_google_calendar(database,client=client,today=today,config_path=config)
    choose_training(sid,'MOVE',database,workout={'date':'2026-11-01','title':'Moved ride','sport':'cycling','duration_seconds':3600})
    result=sync_google_calendar(database,client=client,today=today,config_path=config)
    assert result.updated==1 and result.removed==0 and result.created==0
    assert client.updated[-1][1]['start']['date']=='2026-11-01'


def test_base_limits_optional_extra_quality_when_other_sports_already_costly():
    database=_database();sid=_plan_session(database,'2026-10-02','running','Tempo intervals',2400,intensity='Tempo',priority='C')
    for day in ('2026-09-30','2026-10-01'):_activity(database,day,'cycling',3600,load=150,intensity_factor=1.0)
    with connect(database) as connection:
        pid=connection.execute('SELECT plan_id FROM planned_sessions WHERE id=?',(sid,)).fetchone()[0]
        wid=connection.execute("INSERT INTO training_weeks(plan_id,week_start,week_end,phase_name,raw_json) VALUES (?,'2026-09-28','2026-10-04','Base','{}')",(pid,)).lastrowid
        connection.execute('UPDATE planned_sessions SET week_id=? WHERE id=?',(wid,sid));connection.commit()
    result=adjust_week(date(2026,10,2),database,as_of=date(2026,10,2),persist=False)[0]
    assert result.action=='REDUCE_INTENSITY' and 'base_limit_optional_intensity_accumulation' in result.rules_triggered


def test_athlete_control_ui_can_skip_and_clear_without_overwriting_recommendation(monkeypatch):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    from dashboard import plan_controls
    database=_database();sid=_plan_session(database,'2026-10-02','cycling','Easy Z2',3600)
    adjust_week(date(2026,10,1),database,as_of=date(2026,10,1))
    monkeypatch.setenv('ENDURANCE_DB_PATH',str(database.resolve()))
    monkeypatch.setattr(plan_controls,'_refresh',lambda:None)
    app=AppTest.from_file(str(Path(__file__).parent/'fixtures'/'plan_controls_app.py')).run()
    assert not app.exception
    app.selectbox[0].select('Skip').run()
    next(b for b in app.button if b.label=='Save decision').click().run()
    assert not app.exception and training_decisions(database)[sid]['action']=='SKIP'
    app.selectbox[0].select('Clear decision').run()
    next(b for b in app.button if b.label=='Save decision').click().run()
    assert not app.exception and sid not in training_decisions(database)


@pytest.mark.parametrize('invalid',[False,True])
def test_normal_sync_detects_revision_and_preserves_plan_on_failure(invalid,monkeypatch):
    from tests.test_sync import DiscoveryTransport, FakeResponse, _page, _activity as sync_activity
    from endurance_lab.sync_pipeline import SyncPipeline, SyncOptions
    from endurance_lab.strava_manifest import save_manifest
    from endurance_lab import plan_importer
    root,workbook,database=fixture()
    before=revision_status(database)['revision']['id']
    if invalid:
        copy=load_workbook(workbook);copy['Training Calendar'].title='Invalid';copy.save(workbook);copy.close()
    else:edit(workbook,'K4',120)
    parse_calls=[];original=plan_importer.parse_workbook
    def parse(value):
        parse_calls.append(value);return original(value)
    monkeypatch.setattr(plan_importer,'parse_workbook',parse)
    manifest=root/'manifest.csv';save_manifest([],manifest)
    activity=sync_activity('10000001','2026-09-24T08:00:00+0000',sport='WeightTraining')
    transport=lambda *args,**kwargs:DiscoveryTransport([_page([activity],1,1)],[FakeResponse(404),FakeResponse(404)])
    auth=lambda context:type('Auth',(),{'authenticated':True,'reason':'test'})()
    pipeline=SyncPipeline(transport_factory=transport,authentication_checker=auth)
    options=SyncOptions(database=database,manifest=manifest,downloads=root/'downloads',skip_calendar=True,delay_seconds=0,force_discovery=True,date_from=date(2026,9,24),date_to=date(2026,9,24))
    result=pipeline.run(options)
    assert result.status=='success'
    if invalid:
        assert result.reconciliation['plan_revision']['status']=='failed'
        assert result.warnings and revision_status(database)['revision']['id']==before
    else:
        assert result.reconciliation['plan_revision']['status']=='imported'
        assert revision_status(database)['revision']['id']>before
        again=pipeline.run(options)
        assert again.status=='success' and again.reconciliation['plan_revision']['status']=='unchanged'
        assert len(parse_calls)==1
