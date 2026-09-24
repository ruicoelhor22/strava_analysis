from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from endurance_lab.config import project_path
from endurance_lab.plan_models import (
    ParsedTrainingPlan,
    PlannedSession,
    PlanTarget,
    TrainingPhase,
    TrainingWeek,
)


CALENDAR_SHEET = "Training Calendar"
REQUIRED_CALENDAR_HEADERS = {"Date", "Sport", "Planned Session"}
ACTUAL_FIELDS = (
    "Actual Sport", "Strava Activity ID", "Actual Session", "Actual Duration (min)",
    "Distance (km)", "Elevation (m)", "Avg HR", "Max HR", "Avg Power",
    "Weighted Power", "RPE", "Estimated Session Load", "Completion %",
    "Coach Analysis / Notes", "Athlete Notes", "Constraints / Changes",
)


def discover_workbook(value: str | Path | None = None) -> Path:
    if value is not None:
        resolved = project_path(value).resolve()
        if not resolved.is_file():
            raise FileNotFoundError(resolved)
        if resolved.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError(f"Training plan must be an .xlsx or .xlsm workbook: {resolved}")
        return resolved
    root = project_path("training_data")
    candidates = sorted(
        path.resolve()
        for path in root.glob("*.xls*")
        if path.is_file() and not path.name.startswith("~$") and path.suffix.lower() in {".xlsx", ".xlsm"}
    )
    if len(candidates) == 1:
        return candidates[0]
    preferred = [
        path for path in candidates
        if any(token in path.stem.lower() for token in ("coach", "training", "plan"))
    ]
    if len(preferred) == 1:
        return preferred[0]
    names = ", ".join(path.name for path in candidates) or "none"
    raise ValueError(f"Could not choose one coaching workbook under training_data (found: {names})")


def inspect_workbook(value: str | Path | None = None) -> dict[str, Any]:
    source = discover_workbook(value)
    workbook = load_workbook(source, data_only=False, read_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        formulas = sum(
            1 for row in sheet.iter_rows() for cell in row if cell.data_type == "f"
        )
        nonempty = sum(
            1 for row in sheet.iter_rows() for cell in row if cell.value not in (None, "")
        )
        sheets.append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "dimensions": sheet.calculate_dimension(),
                "nonempty_cells": nonempty,
                "formulas": formulas,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
            }
        )
    calendar = workbook[CALENDAR_SHEET] if CALENDAR_SHEET in workbook.sheetnames else None
    header_row = _find_header_row(calendar) if calendar else None
    return {
        "workbook": str(source),
        "size_bytes": source.stat().st_size,
        "modified_at": datetime.fromtimestamp(source.stat().st_mtime).isoformat(),
        "sheets": sheets,
        "calendar_header_row": header_row,
        "calendar_headers": _headers(calendar, header_row) if calendar and header_row else [],
    }


def parse_workbook(value: str | Path | None = None) -> ParsedTrainingPlan:
    source = discover_workbook(value)
    values_book = load_workbook(source, data_only=True, read_only=False)
    formulas_book = load_workbook(source, data_only=False, read_only=False)
    if CALENDAR_SHEET not in values_book.sheetnames:
        raise ValueError(f"Workbook has no {CALENDAR_SHEET!r} sheet")
    calendar = values_book[CALENDAR_SHEET]
    header_row = _find_header_row(calendar)
    if header_row is None:
        raise ValueError("Training Calendar headers could not be located")
    headers = _headers(calendar, header_row)
    missing = REQUIRED_CALENDAR_HEADERS - set(headers)
    if missing:
        raise ValueError(f"Training Calendar is missing required columns: {sorted(missing)}")

    dated_rows: list[tuple[int, dict[str, Any]]] = []
    for row_number in range(header_row + 1, calendar.max_row + 1):
        row = {
            header: calendar.cell(row_number, index + 1).value
            for index, header in enumerate(headers)
        }
        parsed_date = _date(row.get("Date"))
        if parsed_date is not None:
            row["Date"] = parsed_date
            row["Week Start"] = _date(row.get("Week Start")) or _monday(parsed_date)
            dated_rows.append((row_number, row))
    if not dated_rows:
        raise ValueError("Training Calendar contains no dated rows")

    sessions = _sessions(dated_rows)
    weeks = _weeks(dated_rows)
    phases = _phases(dated_rows)
    snapshot = _workbook_snapshot(values_book, formulas_book)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    return ParsedTrainingPlan(
        source=source,
        name=source.stem.replace("_", " "),
        source_sha256=digest,
        date_start=min(row[1]["Date"] for row in dated_rows),
        date_end=max(row[1]["Date"] for row in dated_rows),
        sheets=tuple(values_book.sheetnames),
        phases=tuple(phases),
        weeks=tuple(weeks),
        sessions=tuple(sessions),
        workbook_snapshot=snapshot,
    )


def _sessions(rows: list[tuple[int, dict[str, Any]]]) -> list[PlannedSession]:
    result: list[PlannedSession] = []
    keys: Counter[str] = Counter()
    for row_number, row in rows:
        title = _text(row.get("Planned Session"))
        sport_raw = _text(row.get("Sport"))
        if not title or not sport_raw:
            continue
        planned_date = row["Date"]
        sport = normalize_plan_sport(sport_raw)
        workout_code = _text(row.get("Workout Code"))
        base_key = f"{planned_date.isoformat()}:{sport}:{_slug(workout_code or title)}"
        keys[base_key] += 1
        source_key = base_key if keys[base_key] == 1 else f"{base_key}:{keys[base_key]}"
        description = _text(row.get("Detailed Prescription"))
        duration = _number(row.get("Planned Duration (min)"))
        priority = _text(row.get("Priority"))
        actual = {key: _json_value(row.get(key)) for key in ACTUAL_FIELDS if row.get(key) not in (None, "")}
        targets = tuple(
            target for target in (
                _target("power", row.get("Power Target"), "W"),
                _target("heart_rate", row.get("HR Target"), "bpm"),
                _target("pace", row.get("Pace Target"), None),
            ) if target is not None
        )
        raw = {key: _json_value(value) for key, value in row.items() if value not in (None, "")}
        result.append(
            PlannedSession(
                source_key=source_key,
                source_row=row_number,
                planned_date=planned_date,
                week_start=row["Week Start"],
                phase=_text(row.get("Phase")) or "Unspecified",
                week_type=_text(row.get("Week Type")),
                sport=sport,
                session_type=_session_type(title, description, row.get("Planned Intensity")),
                workout_code=workout_code,
                title=title,
                description=description,
                planned_duration_seconds=duration * 60 if duration is not None else None,
                planned_distance_m=_planned_distance(title),
                intensity=_text(row.get("Planned Intensity")),
                interval_structure=description,
                priority=priority,
                notes=" | ".join(
                    part for part in (
                        _text(row.get("Session Purpose")),
                        _text(row.get("Fuel / Hydration")),
                        _text(row.get("Gym Details")),
                    ) if part
                ) or None,
                source_status=_text(row.get("Status")),
                workbook_actual=actual,
                targets=targets,
                raw_source=raw,
            )
        )
    return result


def _weeks(rows: list[tuple[int, dict[str, Any]]]) -> list[TrainingWeek]:
    grouped: dict[date, list[dict[str, Any]]] = defaultdict(list)
    for _, row in rows:
        grouped[row["Week Start"]].append(row)
    result = []
    for week_start, items in sorted(grouped.items()):
        phases = Counter(_text(item.get("Phase")) or "Unspecified" for item in items)
        week_types = Counter(_text(item.get("Week Type")) for item in items if _text(item.get("Week Type")))
        purposes = list(dict.fromkeys(
            _text(item.get("Session Purpose")) for item in items if _text(item.get("Session Purpose"))
        ))
        result.append(
            TrainingWeek(
                week_start=week_start,
                week_end=week_start + timedelta(days=6),
                phase=phases.most_common(1)[0][0],
                week_type=week_types.most_common(1)[0][0] if week_types else None,
                objectives="; ".join(purposes) or None,
                raw_source={"dates": [item["Date"].isoformat() for item in items]},
            )
        )
    return result


def _phases(rows: list[tuple[int, dict[str, Any]]]) -> list[TrainingPhase]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for _, row in rows:
        grouped[_text(row.get("Phase")) or "Unspecified"].append(row)
    result = []
    for name, items in grouped.items():
        purposes = list(dict.fromkeys(
            _text(item.get("Session Purpose")) for item in items if _text(item.get("Session Purpose"))
        ))
        result.append(
            TrainingPhase(
                name=name,
                start_date=min(item["Date"] for item in items),
                end_date=max(item["Date"] for item in items),
                objectives="; ".join(purposes) or None,
                raw_source={"calendar_rows": len(items)},
            )
        )
    return sorted(result, key=lambda item: item.start_date)


def normalize_plan_sport(value: str | None) -> str:
    text = _text(value).lower().replace(" ", "")
    if text in {"ride", "cycling", "bike"}:
        return "cycling"
    if text in {"run", "running"}:
        return "running"
    if text in {"weighttraining", "strength", "gym"}:
        return "strength"
    if text in {"swim", "swimming"}:
        return "swimming"
    if text in {"hike", "hiking"}:
        return "hiking"
    if text in {"walk", "walking"}:
        return "walking"
    return "other"


def _workbook_snapshot(values_book, formulas_book) -> dict[str, Any]:
    sheets: dict[str, Any] = {}
    for sheet in values_book.worksheets:
        formula_sheet = formulas_book[sheet.title]
        rows = []
        for row in sheet.iter_rows():
            cells = []
            for cell in row:
                if cell.value in (None, ""):
                    continue
                formula = formula_sheet[cell.coordinate].value
                cells.append(
                    {
                        "cell": cell.coordinate,
                        "value": _json_value(cell.value),
                        "formula": formula if isinstance(formula, str) and formula.startswith("=") else None,
                        "style_id": cell.style_id,
                        "number_format": cell.number_format,
                    }
                )
            if cells:
                rows.append(cells)
        sheets[sheet.title] = {
            "dimensions": sheet.calculate_dimension(),
            "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
            "rows": rows,
        }
    return {"sheets": sheets}


def _find_header_row(sheet) -> int | None:
    if sheet is None:
        return None
    for row in range(1, min(sheet.max_row, 20) + 1):
        values = {_text(sheet.cell(row, column).value) for column in range(1, sheet.max_column + 1)}
        if REQUIRED_CALENDAR_HEADERS.issubset(values):
            return row
    return None


def _headers(sheet, row: int | None) -> list[str]:
    if sheet is None or row is None:
        return []
    return [
        _text(sheet.cell(row, column).value) or f"column_{get_column_letter(column)}"
        for column in range(1, sheet.max_column + 1)
    ]


def _target(target_type: str, value: Any, unit: str | None) -> PlanTarget | None:
    raw = _text(value)
    if not raw:
        return None
    numbers = [float(item.replace(",", ".")) for item in re.findall(r"(?<![A-Za-z])\d+(?:[.,]\d+)?", raw)]
    minimum = maximum = None
    confidence = "low"
    if target_type in {"power", "heart_rate"} and numbers:
        relevant = [number for number in numbers if number > (40 if target_type == "heart_rate" else 50)]
        if len(relevant) == 1:
            minimum = maximum = relevant[0]
            confidence = "medium"
        elif len(relevant) == 2:
            minimum, maximum = min(relevant), max(relevant)
            confidence = "high"
    return PlanTarget(target_type, raw, minimum, maximum, unit, confidence)


def _planned_distance(title: str) -> float | None:
    match = re.search(r"(?:~|≈)?\s*(\d+(?:[.,]\d+)?)\s*km\b", title, flags=re.IGNORECASE)
    return float(match.group(1).replace(",", ".")) * 1000 if match else None


def _session_type(*values: Any) -> str | None:
    text = " ".join(_text(value).lower() for value in values if value)
    for key, label in (
        ("threshold", "threshold"), ("vo2", "vo2"), ("interval", "intervals"),
        ("long", "long"), ("recovery", "recovery"), ("easy", "easy"),
        ("z2", "aerobic"), ("aerobic", "aerobic"), ("gym", "strength"),
        ("strength", "strength"), ("race", "race"),
    ):
        if key in text:
            return label
    return None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip()).date()
        except ValueError:
            return None
    return None


def _monday(value: date) -> date:
    return value - timedelta(days=value.weekday())


def _number(value: Any) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")[:80] or "session"


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)
